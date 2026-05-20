"""
Sample Document Generator

=== INTERVIEW EXPLANATION ===
This script generates sample test documents so the project is fully 
self-contained and testable out of the box.

It creates:
1. samples/sample_structure_good.docx:
   - A well-structured document with Heading 1/2 styles, normal paragraphs,
     and an embedded table. Should return a HIGH score.
     
2. samples/sample_structure_poor.docx:
   - A document with long, dense paragraphs, no heading styles,
     and repetitious content. Should return a LOW score.
"""

import os
from docx import Document


def create_good_sample(output_path: str):
    """Creates a well-structured Word document with proper headings & table."""
    doc = Document()
    
    doc.add_heading("API Integration & Authentication Guide", level=0)
    
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Welcome to the API Integration Guide. This document provides step-by-step "
        "instructions on how to authenticate, retrieve data, and handle errors using "
        "our REST API services. Please review the security protocols before writing code."
    )
    
    doc.add_heading("2. Authentication Protocol", level=1)
    doc.add_paragraph(
        "All requests to the API must include a valid bearer token in the HTTP Authorization header. "
        "To obtain an API key, log in to your developer console, navigate to the Credentials tab, "
        "and click 'Generate New Token'."
    )
    
    doc.add_heading("2.1 Bearer Token Header Example", level=2)
    doc.add_paragraph(
        "The header must be formatted as follows:\n"
        "Authorization: Bearer YOUR_API_KEY_HERE"
    )
    
    doc.add_heading("3. Supported Endpoints", level=1)
    doc.add_paragraph(
        "The following endpoints are currently active for the v1 release of the services API:"
    )
    
    # Add a table
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Method'
    hdr_cells[1].text = 'Endpoint'
    hdr_cells[2].text = 'Description'
    
    row1 = table.rows[1].cells
    row1[0].text = 'GET'
    row1[1].text = '/v1/users'
    row1[2].text = 'Retrieve user profiles'
    
    row2 = table.rows[2].cells
    row2[0].text = 'POST'
    row2[1].text = '/v1/orders'
    row2[2].text = 'Submit a transaction'

    doc.add_paragraph("") # Spacing
    doc.add_heading("4. Error Handling", level=1)
    doc.add_paragraph(
        "Standard HTTP status codes are returned to indicate the success or failure of an API request. "
        "Codes in the 2xx range indicate success, 4xx range indicate client errors, "
        "and 5xx range indicate server errors. Ensure your application includes try-catch blocks."
    )
    
    doc.save(output_path)
    print(f"Good sample created at: {output_path}")


def create_poor_sample(output_path: str):
    """Creates a poorly structured Word document with no headings & huge paragraphs."""
    doc = Document()
    
    # Title is normal text instead of Heading style
    p_title = doc.add_paragraph()
    run = p_title.add_run("unstructured_document_final_copy_v2_draft")
    run.bold = True
    
    # Massive, dense paragraph 1
    doc.add_paragraph(
        "This is the document that explains how to do stuff in the system but it is not "
        "very well structured because it does not have any heading styles or clear divisions. "
        "We are writing this to test the analyzer tool. You need to make sure the server is running "
        "before you try to upload anything and you also need to make sure you have the python environment "
        "fully activated with all libraries installed. First you go to the dashboard then you click on "
        "the user profile page and you click the settings button to update your profile picture. If the "
        "settings button is not showing up you should check if your internet is working or reload the page "
        "multiple times. The profile picture must be less than two megabytes in size otherwise it will show "
        "an error and fail to upload. To update your password you go to the security settings page and type "
        "your current password followed by your new password twice. Remember to make your password complex "
        "with symbols and numbers to keep it secure from hackers. If you get locked out of your account "
        "you should contact support at admin@example.com and they will send you a reset link via email."
    )
    
    # Massive, dense paragraph 2
    doc.add_paragraph(
        "Now we will talk about the other modules in the application which includes the database backup "
        "scheduler which runs every midnight to copy database records and files to an offsite secure storage "
        "bucket. The scheduler is configured using a cron job that runs a python backup script. The backup script "
        "first dumps the PostgreSQL database into a SQL file, compresses the SQL file into a GZIP archive, and "
        "uploads it to AWS S3. The AWS S3 bucket has a lifecycle policy that automatically deletes backups older "
        "than thirty days to save storage costs. If you need to restore a backup you must download the compressed "
        "file from S3, unzip it, and run the psql restore command on your server database. You should test the "
        "restore process once every three months on a staging environment to ensure the backup files are not "
        "corrupted and can be restored successfully."
    )
    
    doc.save(output_path)
    print(f"Poor sample created at: {output_path}")


if __name__ == "__main__":
    samples_dir = "samples"
    os.makedirs(samples_dir, exist_ok=True)
    
    create_good_sample(os.path.join(samples_dir, "sample_structure_good.docx"))
    create_poor_sample(os.path.join(samples_dir, "sample_structure_poor.docx"))
