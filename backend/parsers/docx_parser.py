"""
DOCX Parser Module
Extracts text, structure, and metadata from Microsoft Word (.docx) files.

=== INTERVIEW EXPLANATION ===
How .docx files work:
- A .docx file is NOT a plain text file - it's actually a ZIP archive!
- Inside, it contains XML files that define the document structure.
- The python-docx library reads these XML files and gives us a Python API.

What we extract:
- Paragraphs: Each block of text in the document
- Headings: Paragraphs with styles like "Heading 1", "Heading 2", etc.
- Tables: Whether the document contains any tables
- Images: Whether the document has embedded images

Why this matters for migration:
- Headings define the document structure (table of contents)
- Without headings, a document is just a wall of text — hard to migrate
- Tables and images may need special handling in Document360
"""

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.text.paragraph import Paragraph
from docx.table import Table
from backend.models import ParsedDocument, ParagraphMetadata
import os
import math


def _table_to_markdown(table) -> str:
    """Helper to convert python-docx Table object to GFM Markdown table syntax."""
    markdown_lines = []
    rows = table.rows
    if not rows:
        return ""
    
    # Extract header cells
    header_cells = rows[0].cells
    header_text = [cell.text.strip().replace("\n", " ") for cell in header_cells]
    markdown_lines.append("| " + " | ".join(header_text) + " |")
    
    # Add markdown separator line
    markdown_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")
    
    # Add data rows
    for row in rows[1:]:
        row_cells = row.cells
        row_text = [cell.text.strip().replace("\n", " ") for cell in row_cells]
        markdown_lines.append("| " + " | ".join(row_text) + " |")
        
    return "\n".join(markdown_lines)


def parse_docx(file_path: str) -> ParsedDocument:
    """
    Parse a .docx file and extract structured content.
    
    Args:
        file_path: Path to the .docx file
        
    Returns:
        ParsedDocument with extracted text, headings, paragraphs, etc.
    """
    file_size_kb = round(os.path.getsize(file_path) / 1024, 2)
    file_name = os.path.basename(file_path)
    
    try:
        doc = Document(file_path)
    except PackageNotFoundError:
        return ParsedDocument(
            file_name=file_name,
            file_type=".docx",
            file_size_kb=file_size_kb,
            full_text="[Error: Could not open document. File may be corrupted.]",
            paragraphs=[],
            headings=[],
            total_pages=0,
        )
    except Exception as e:
        return ParsedDocument(
            file_name=file_name,
            file_type=".docx",
            file_size_kb=file_size_kb,
            full_text=f"[Error: {str(e)}]",
            paragraphs=[],
            headings=[],
            total_pages=0,
        )
    
    paragraphs = []
    headings = []
    full_text_parts = []
    markdown_parts = []
    paragraph_metadata = []
    
    # --- TABLE DETECTION ---
    has_tables = len(doc.tables) > 0
    
    # --- IMAGE DETECTION ---
    has_images = False
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            has_images = True
            break
            
    # --- MAIN PARSING LOOP (Order Preserving) ---
    # Iterate over child elements of the body to keep text & tables in correct order
    cumulative_words = 0
    for child in doc.element.body:
        # Paragraph element
        if child.tag.endswith('p'):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
                
            words_in_para = len(text.split())
            cumulative_words += words_in_para
            full_text_parts.append(text)
            style_name = para.style.name if para.style else ""
            
            # Estimate page dynamically based on standard 350 words per page
            page_num = max(1, math.ceil(cumulative_words / 350))
            paragraph_metadata.append(ParagraphMetadata(
                text=text,
                page=page_num,
                line=0
            ))
            
            if style_name.startswith("Heading") or style_name == "Title":
                # Map styles to Heading levels
                level = 1
                if "1" in style_name:
                    level = 1
                elif "2" in style_name:
                    level = 2
                elif "3" in style_name:
                    level = 3
                elif "4" in style_name:
                    level = 4
                
                headings.append(f"H{level}: {text}")
                markdown_parts.append(f"{'#' * level} {text}")
            else:
                paragraphs.append(text)
                markdown_parts.append(text)
                
        # Table element
        elif child.tag.endswith('tbl'):
            table = Table(child, doc)
            table_markdown = _table_to_markdown(table)
            if table_markdown:
                markdown_parts.append(table_markdown)
                # Add table text to full_text for metrics tracking
                table_text = " ".join(cell.text.strip() for row in table.rows for cell in row.cells)
                full_text_parts.append(table_text)
                
                # Count table words to keep page estimation accurate
                words_in_table = len(table_text.split())
                cumulative_words += words_in_table
    
    full_text = "\n\n".join(full_text_parts)
    markdown_content = "\n\n".join(markdown_parts)
    
    # --- PAGE ESTIMATION ---
    estimated_pages = max(1, math.ceil(cumulative_words / 350))
    
    return ParsedDocument(
        file_name=file_name,
        file_type=".docx",
        file_size_kb=file_size_kb,
        full_text=full_text,
        markdown_content=markdown_content,
        paragraphs=paragraphs,
        headings=headings,
        total_pages=estimated_pages,
        has_tables=has_tables,
        has_images=has_images,
        paragraph_metadata=paragraph_metadata,
    )
    

